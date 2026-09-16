"""Сборка .docx из документации в docs/*.md.

Схемы Mermaid подставляются готовыми картинками из docs/assets в порядке
их появления в тексте. Пересобрать картинки:
    npx @mermaid-js/mermaid-cli -i схема.mmd -o docs/assets/arch-N.png -w 1500 -b white -s 2

Запуск:  python scripts/md2docx.py
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"
ASSETS = DOCS / "assets"

BODY_FONT = "Calibri"
MONO_FONT = "Consolas"
ACCENT = RGBColor(0x1F, 0x38, 0x64)
CODE_BG = "F2F3F5"
HEAD_BG = "E7EAF0"

INLINE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\[[^\]]+\]\([^)]+\))")
LINK = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")


# --- низкоуровневые помощники ------------------------------------------------


def shade(element, fill: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    element.append(shd)


def cell_shade(cell, fill: str) -> None:
    shade(cell._tc.get_or_add_tcPr(), fill)


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    hdr = OxmlElement("w:tblHeader")
    hdr.set(qn("w:val"), "true")
    tr_pr.append(hdr)


def bottom_border(paragraph, size: int = 6, color: str = "C9CDD6") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "1F5FA8")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.append(color)
    r_pr.append(underline)
    run.append(r_pr)
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    link.append(run)
    paragraph._p.append(link)


# --- разметка ----------------------------------------------------------------


def add_inline(paragraph, text: str) -> None:
    """Разбирает `код`, **жирный** и ссылки в один абзац."""
    for token in INLINE.split(text):
        if not token:
            continue
        if token.startswith("`") and token.endswith("`") and len(token) > 1:
            run = paragraph.add_run(token[1:-1])
            run.font.name = MONO_FONT
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x8A, 0x1F, 0x4C)
        elif token.startswith("**") and token.endswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif (match := LINK.fullmatch(token)) is not None:
            label, target = match.group(1), match.group(2)
            if target.startswith("http"):
                add_hyperlink(paragraph, label, target)
            else:
                # ссылка на соседний .md — в Word её некуда вести
                run = paragraph.add_run(label)
                run.italic = True
        else:
            paragraph.add_run(token)


def add_code_block(doc: Document, lines: list[str]) -> None:
    text = "\n".join(lines).rstrip()
    paragraph = doc.add_paragraph()
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(6)
    fmt.space_after = Pt(10)
    fmt.left_indent = Cm(0.4)
    shade(paragraph._p.get_or_add_pPr(), CODE_BG)
    run = paragraph.add_run(text)
    run.font.name = MONO_FONT
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x1B, 0x1B, 0x1B)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = True

    for r, row in enumerate(rows):
        for c in range(width):
            cell = table.cell(r, c)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.space_after = Pt(2)
            value = row[c] if c < len(row) else ""
            add_inline(paragraph, value)
            for run in paragraph.runs:
                run.font.size = Pt(9)
                if r == 0:
                    run.bold = True
            if r == 0:
                cell_shade(cell, HEAD_BG)
    repeat_header(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def usable_width(section) -> float:
    return section.page_width - section.left_margin - section.right_margin


def portrait(section) -> None:
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Cm(2))
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)


def landscape(section) -> None:
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Cm(29.7), Cm(21)
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Cm(1.5))
    section.top_margin = Cm(1.5)
    section.bottom_margin = Cm(1.5)


def add_image(doc: Document, path: Path, caption: str | None = None) -> None:
    """Широкие схемы уезжают на отдельную страницу в альбомной ориентации."""
    px_w, px_h = Image.open(path).size
    wide = px_w / px_h > 2.0

    if wide:
        landscape(doc.add_section(WD_SECTION.NEW_PAGE))

    section = doc.sections[-1]
    max_w = usable_width(section)
    max_h = section.page_height - section.top_margin - section.bottom_margin - Cm(1.5)

    width = max_w
    if width * px_h / px_w > max_h:
        width = int(max_h * px_w / px_h)

    holder = doc.add_paragraph()
    holder.alignment = WD_ALIGN_PARAGRAPH.CENTER
    holder.paragraph_format.space_before = Pt(8)
    holder.paragraph_format.space_after = Pt(4)
    holder.add_run().add_picture(str(path), width=width)

    if caption:
        line = doc.add_paragraph()
        line.alignment = WD_ALIGN_PARAGRAPH.CENTER
        line.paragraph_format.space_after = Pt(12)
        run = line.add_run(caption)
        run.font.size = Pt(8.5)
        run.italic = True
        run.font.color.rgb = RGBColor(0x60, 0x66, 0x72)

    if wide:
        portrait(doc.add_section(WD_SECTION.NEW_PAGE))


# --- документ ---------------------------------------------------------------


def new_document() -> Document:
    doc = Document()
    portrait(doc.sections[0])

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, before in (
        ("Heading 1", 17, 20),
        ("Heading 2", 13.5, 16),
        ("Heading 3", 11.5, 12),
        ("Heading 4", 10.5, 10),
    ):
        style = doc.styles[name]
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        style.font.color.rgb = ACCENT
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    title = doc.styles["Title"]
    title.font.name = BODY_FONT
    title.font.size = Pt(26)
    title.font.color.rgb = ACCENT
    title.font.bold = True

    return doc


def parse_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_separator(line: str) -> bool:
    return bool(re.fullmatch(r"\|[\s\-:|]+\|", line.strip()))


def convert(md_path: Path, out_path: Path, images: list[tuple[Path, str]]) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = new_document()
    image_queue = list(images)
    first_heading = True
    i = 0

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # --- код и схемы ---
        if stripped.startswith("```"):
            lang = stripped[3:].strip().lower()
            body: list[str] = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                body.append(lines[i])
                i += 1
            i += 1
            if lang == "mermaid":
                if image_queue:
                    path, caption = image_queue.pop(0)
                    add_image(doc, path, caption)
                else:
                    add_code_block(doc, body)
            else:
                add_code_block(doc, body)
            continue

        # --- горизонтальная линия ---
        if re.fullmatch(r"-{3,}", stripped):
            rule = doc.add_paragraph()
            rule.paragraph_format.space_before = Pt(2)
            rule.paragraph_format.space_after = Pt(2)
            bottom_border(rule)
            i += 1
            continue

        # --- заголовки ---
        if (match := re.match(r"^(#{1,4})\s+(.*)$", stripped)) is not None:
            level, text = len(match.group(1)), match.group(2)
            if level == 1 and first_heading:
                paragraph = doc.add_paragraph(style="Title")
                add_inline(paragraph, text)
                for run in paragraph.runs:
                    run.font.size = Pt(26)
                    run.bold = True
                first_heading = False
            else:
                paragraph = doc.add_paragraph(style=f"Heading {min(level, 4)}")
                add_inline(paragraph, text)
                for run in paragraph.runs:
                    run.font.color.rgb = ACCENT
            i += 1
            continue

        # --- таблицы ---
        if stripped.startswith("|"):
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                current = lines[i].strip()
                if not is_separator(current):
                    rows.append(parse_table_row(current))
                i += 1
            add_table(doc, rows)
            continue

        # --- списки ---
        if (match := re.match(r"^[-*]\s+(.*)$", stripped)) is not None:
            text = match.group(1)
            checkbox = re.match(r"^\[([ xX])\]\s*(.*)$", text)
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(3)
            if checkbox:
                mark = "\u2611 " if checkbox.group(1).lower() == "x" else "\u2610 "
                paragraph.add_run(mark)
                add_inline(paragraph, checkbox.group(2))
            else:
                add_inline(paragraph, text)
            i += 1
            continue

        if (match := re.match(r"^\d+\.\s+(.*)$", stripped)) is not None:
            paragraph = doc.add_paragraph(style="List Number")
            paragraph.paragraph_format.space_after = Pt(3)
            add_inline(paragraph, match.group(1))
            i += 1
            continue

        # --- обычный абзац ---
        paragraph = doc.add_paragraph()
        add_inline(paragraph, stripped)
        i += 1

    doc.core_properties.title = md_path.stem
    doc.core_properties.author = "VTBIH"
    doc.save(out_path)
    print(f"{out_path.name}: готово")


CAPTIONS = [
    "Схема 1. Контекст: внешние системы и связи",
    "Схема 2. Развёртывание: контейнеры, порты и тома",
    "Схема 3. Конвейер обработки документа: два режима",
    "Схема 4. Модель данных",
]

JOBS = [
    ("user-guide.md", "VTBIH-Руководство-пользователя.docx", []),
    ("deployment.md", "VTBIH-Развёртывание.docx", []),
    (
        "architecture.md",
        "VTBIH-Архитектура.docx",
        [(ASSETS / f"arch-{n}.png", CAPTIONS[n - 1]) for n in (1, 2, 3, 4)],
    ),
]


def main() -> int:
    for src, dst, images in JOBS:
        md_path = DOCS / src
        if not md_path.exists():
            print(f"нет файла {md_path}", file=sys.stderr)
            return 1
        for path, _ in images:
            if not path.exists():
                print(f"нет картинки {path}", file=sys.stderr)
                return 1
        convert(md_path, DOCS / dst, images)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
